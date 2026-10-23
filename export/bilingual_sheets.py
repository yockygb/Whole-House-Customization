"""English-first auxiliary worksheets linked to the owning calculation sheet."""
from openpyxl.styles import Font,Alignment,PatternFill,Border,Side
from core.floors import floor_name_en

BLUE='0B3D66';GREY='748094'
def bi(en,cn):return str(en or 'English translation needed')+'\n'+str(cn or '')
def style(ws):
    ws.sheet_view.showGridLines=False
    for row in ws:
        for c in row:
            c.font=Font(name='Arial',size=10,color=BLUE if c.column==1 else '243248')
            c.alignment=Alignment(vertical='center',wrap_text=True)
    ws.freeze_panes='B4';ws.page_setup.orientation='landscape';ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0

def summary_sheet(ws,cfg,lines,summary,kind='internal'):
    ws['A1']=bi('Summary by group','分类汇总');ws.merge_cells('A1:E1');ws.row_dimensions[1].height=35
    ws['A2']=bi('Values link to calculation details and update when dimensions or rates change.','数值引用计算明细，修改尺寸或价格后联动。');ws.merge_cells('A2:E2');ws.row_dimensions[2].height=35
    show=kind=='internal';r=4
    for en,cn,key in [('By floor','按楼层','floor'),('By room','按区域','room'),('By category','按品类','cat_name'),('By material','按材料','material_label')]:
        ws.cell(r,1,bi(en,cn));r+=1
        heads=[bi('Name','名称'),bi('Rows','行数'),bi('Amount (RM)','金额')]+([bi('Cost (RM)','成本'),bi('Gross profit (RM)','毛利')] if show else [])
        for col,h in enumerate(heads,1):ws.cell(r,col,h).fill=PatternFill('solid',fgColor='EEF3F8')
        ws.row_dimensions[r].height=30;r+=1;start=r
        groups={}
        for i,ln in enumerate(lines):groups.setdefault(getattr(ln,key),[]).append((i,ln))
        for label,members in groups.items():
            first=members[0][1]
            english={'floor':first.floor_en or floor_name_en(label),'room':first.room_en,'cat_name':first.cat_name_en,'material_label':first.material_label_en}[key]
            ws.cell(r,1,bi(english,label));ws.cell(r,2,len(members))
            refs=','.join('明细_Calculation!V'+str(5+i) for i,_ in members)
            ws.cell(r,3,'=ROUND(SUM('+refs+'),2)').number_format='#,##0.00'
            if show:
                refs=','.join('成本毛利_CostProfit!H'+str(4+i) for i,_ in members)
                ws.cell(r,4,'=ROUND(SUM('+refs+'),2)').number_format='#,##0.00'
                complete=all(getattr(ln,'cost_complete',True) for _,ln in members)
                ws.cell(r,5,f'=C{r}-D{r}' if complete else 'Missing cost / 成本待补').number_format='#,##0.00'
            ws.row_dimensions[r].height=34;r+=1
        ws.cell(r,1,bi('Subtotal','小计'))
        for c in range(2,6 if show else 4):
            from openpyxl.utils import get_column_letter
            letter=get_column_letter(c)
            ws.cell(r,c,f'=SUM({letter}{start}:{letter}{r-1})' if r>start else '=0').number_format='#,##0.00'
        if show and summary.get('cost_complete') is False:ws.cell(r,5,'Missing cost / 成本待补')
        r+=2
    for c,w in {'A':48,'B':10,'C':20,'D':20,'E':20}.items():ws.column_dimensions[c].width=w
    style(ws)

def usage_sheet(ws,cfg,lines):
    ws['A1']=bi('Material usage','材料用量');ws.merge_cells('A1:F1');ws.row_dimensions[1].height=35
    ws['A2']=bi('Pricing measure for procurement; does not include cutting waste or a panel cutting list.','用于备料的计价用量，不含损耗；不等同于板材开料清单。');ws.merge_cells('A2:F2');ws.row_dimensions[2].height=35
    for c,t in enumerate([bi('Material','材料'),bi('Category','品类'),bi('Unit','单位'),bi('Total measure','用量合计'),bi('Rows','行数'),bi('Note','备注')],1):ws.cell(3,c,t)
    groups={}
    for i,l in enumerate(lines):
        mixed=getattr(l,'pricing_model','')=='kitchen_split'
        groups.setdefault((l.body_label if mixed else l.material_label,l.cat_key,l.unit,'body'),[]).append((i,l))
        if mixed and l.door_key:groups.setdefault((l.door_label,l.cat_key,'ft²','door'),[]).append((i,l))
    for r,((label,_,unit,part),members) in enumerate(groups.items(),4):
        first=members[0][1];mixed=getattr(first,'pricing_model','')=='kitchen_split'
        english=first.door_label_en if part=='door' else first.body_label_en if mixed else first.material_label_en
        ws.cell(r,1,bi(english,label));ws.cell(r,2,bi(first.cat_name_en,first.cat_name));ws.cell(r,3,unit)
        a,b=('AC','AD') if part=='door' else ('P','R')
        ws.cell(r,4,'=ROUND('+ '+'.join(f'明细_Calculation!{a}{5+i}*明细_Calculation!{b}{5+i}*明细_Calculation!T{5+i}' for i,_ in members)+',3)').number_format='#,##0.000'
        ws.cell(r,5,len(members));ws.row_dimensions[r].height=38
    for c,w in {'A':50,'B':25,'C':10,'D':20,'E':10,'F':30}.items():ws.column_dimensions[c].width=w
    style(ws)

def notes_sheet(ws,cfg,quote,summary,kind='internal'):
    r=1
    def line(en,cn,heading=False):
        nonlocal r
        ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=6)
        c=ws.cell(r,1,bi(en,cn));c.font=Font(name='Arial',size=11 if heading else 9,bold=heading,color=BLUE)
        c.alignment=Alignment(wrap_text=True,vertical='center');ws.row_dimensions[r].height=34 if heading else 42;r+=1
    line('Quotation notes','报价说明',True)
    line('Pricing rules','计价规则',True)
    line('Area items: length (ft) × billable height (ft) × depth factor × unit price × quantity. Optional door rates are added to the carcass rate.','面积类：长(ft) × 计价高(ft) × 深度系数 × 单价 × 数量；柜门单价另加到柜体单价。')
    line('Length items: length (ft) × non-standard factor × unit price × quantity.','长度类：长(ft) × 非标系数 × 单价 × 数量。')
    if any(c.get('pricing_model')=='kitchen_split' for c in cfg.cat_list()):
        line('Kitchen: carcass length (configured unit: ft; historical metre rates remain readable) × carcass rate × non-standard factor × quantity + doors using ordinary cabinet projection rules (ft²) + add-ons. Base cabinet plinth follows total length; handles are optional.','橱柜：柜体长度(按配置单位，新报价FT，历史米价保留) × 柜体单价 × 非标系数 × 数量＋柜门按普通柜投影规则(ft²)单独计算＋增项。地柜脚线跟随总长度自动计算，拉手选配。')
    for c in cfg.cat_list():
        line(f"{c.get('name_en',c['key'])}: depth {c.get('std_depth_mm',0)} mm; minimum billable height {c.get('min_height_ft',0)} ft.",f"{c.get('name',c['key'])}：标准深度 {c.get('std_depth_mm',0)} mm；最低计价高 {c.get('min_height_ft',0)} ft。")
    line(f'Non-standard multiplier: {cfg.nonstd_factor}; height tolerance: {cfg.height_tol} ft; 1 ft = {cfg.mm_per_ft} mm.',f'非标系数：{cfg.nonstd_factor}；高度容差：{cfg.height_tol} ft；1 ft = {cfg.mm_per_ft} mm。')
    line('Add-ons use the configured length, area or quantity basis. Glass uses its own actual dimensions, not the minimum cabinet height.','增项按配置的长度、面积或数量计；玻璃按单独实际尺寸，不套用柜体最低高度。')
    line('Materials & rates','材料与单价',True)
    for cat in cfg.cat_list():
        for key,e in cfg.prices.get(cat['key'],{}).items():
            en=f"{cfg.cat_name_en(cat['key'])}: {e.get('label_en','English translation needed')} — {cfg.currency} {cfg.unit_price(cat['key'],key,quote.get('mode')):.2f}/{e.get('unit','')}"
            cn=f"{cat.get('name','')}：{e.get('label',key)}，报价 {cfg.unit_price(cat['key'],key,quote.get('mode')):.2f}/{e.get('unit','')}"
            if kind=='internal':
                en+=f"; supply {cfg.supply_price(cat['key'],key):.2f}";cn+=f"；供货价 {cfg.supply_price(cat['key'],key):.2f}"
            line(en,cn)
    line('Quotation information','报价信息',True)
    line(f"Customer: {quote.get('customer','')} | Project: {quote.get('project','')}",f"客户：{quote.get('customer','')}；项目：{quote.get('project','')}")
    line(f"Reference: {quote.get('quote_no','')} | Date: {quote.get('quote_date','')} | Validity: {quote.get('valid_days',30)} days",f"编号：{quote.get('quote_no','')}；日期：{quote.get('quote_date','')}；有效期：{quote.get('valid_days',30)}天")
    line('Exclusions','未含项目',True)
    line('Tax, delivery, installation and hardware upgrades are excluded unless separately listed.','除非另行列明，报价不含税、运输、安装和五金升级。')
    line('Quartz cut-outs, polishing and backsplash are charged separately. Final measurements must be confirmed on site.','石英石开孔、磨边、挡水另计，最终尺寸须现场复核。')
    line('Editable Excel formulas recalculate amounts. Cabinet pictures are embedded static previews; re-export after dimension changes to update pictures.','Excel金额保留可编辑公式并可重算。柜体图片为嵌入式静态图；修改尺寸后请重新导出以更新图片。')
    line('Whole-House Quotation Studio 1.1.','全屋定制报价系统1.1。')
    for c in 'ABCDEF':ws.column_dimensions[c].width=21
    ws.page_setup.orientation='portrait';ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
