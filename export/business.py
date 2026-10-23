"""Formula-linked v5 add-ons and project cost extension; customer layout is reused."""
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def finalize(wb,cfg,quote,lines,summary,kind,price_rows,addon_rows,sum_rows):
    if 'cost_policy' not in summary: return
    from export.xlsx import SHEET_D, SHEET_P, SHEET_C
    detail=wb[SHEET_D]
    addon=wb.create_sheet('增项明细_Addons')
    heads=['Line / 柜体行','Add-on / 增项','Unit / 单位','Qty / 数量','Width mm / 宽',
           'Height mm / 高','Length mm / 长','Measure / 计费量','Retail / 零售价','Amount / 金额']
    if kind=='internal': heads+=['Cost rate / 成本单价','Cost / 成本金额']
    addon.append(heads)
    cost=wb[SHEET_C] if kind=='internal' else None
    r=2
    for i,ln in enumerate(lines):
        dr=5+i
        cat=cfg.category(ln.cat_key)
        if cost is not None:
            cr=4+i
            cost.cell(cr,13,f'={SHEET_D}!I{dr}*{SHEET_D}!J{dr}*{SHEET_D}!T{dr}/1000000'
                      if cat.get('cost_projection',cat['key']!='wall_panel') else 0)
            cost.cell(cr,14,f'=M{cr}*项目成本_ProjectCosts!$B$5')
            cost.cell(cr,15,'Complete / 完整' if ln.cost_complete else 'Missing cost / 成本待补')
            if not ln.cost_complete:
                for col in (10,11):cost.cell(cr,col,'Missing cost / 成本待补')
        if ln.pricing_model=='historical':
            if ln.led_amount:detail.cell(dr,22,detail.cell(dr,22).value+f'+AA{dr}')
            continue
        refs=[]; costrefs=[]
        for e in ln.extras_list:
            row=r;r+=1
            unit=e['unit'].replace('²','2')
            addon.cell(row,1,ln.idx);addon.cell(row,2,(e.get('name_en') or '')+' / '+e['name'])
            addon.cell(row,3,e['unit']);addon.cell(row,4,e['qty'])
            if e.get('auto_with_cabinet'): addon.cell(row,4,f'={SHEET_D}!T{dr}')
            addon.cell(row,5,e.get('w_mm') or f'={SHEET_D}!I{dr}')
            addon.cell(row,6,e.get('h_mm') or f'={SHEET_D}!J{dr}')
            addon.cell(row,7,(f'=IF(SUM({SHEET_D}!AB{dr})>0,SUM({SHEET_D}!AB{dr}),{SHEET_D}!I{dr})'
                              if e['key']=='quartz' else e.get('l_mm') or f'={SHEET_D}!I{dr}'))
            formula=f'=D{row}'
            if e['basis']=='length': formula=f'=G{row}/{1000 if unit=="m" else 304.8}*D{row}'
            if e['basis']=='area': formula=f'=E{row}*F{row}/{1000000 if unit=="m2" else 92903.04}*D{row}'
            addon.cell(row,8,formula)
            ar=addon_rows.get(e['key']) if not cat.get('source_key') else None
            addon.cell(row,9,f'={SHEET_P}!H{ar}' if ar else e['price'])
            addon.cell(row,10,f'=ROUND(H{row}*I{row},2)')
            refs.append(f'增项明细_Addons!J{row}')
            if cost is not None:
                addon.cell(row,11,e.get('cost_price'))
                addon.cell(row,12,f'=ROUND(H{row}*SUM(K{row}),2)')
                costrefs.append(f'增项明细_Addons!L{row}')
        detail.cell(dr,21,'=SUM('+','.join(refs)+')' if refs else 0)
        # Per-category nonstandard rules must agree with the server, not global defaults.
        cat=cfg.category(ln.cat_key)
        mixed=cat.get('pricing_model')=='kitchen_split'
        if cat['measure']=='length':
            factor=cat.get('nonstd_factor',cfg.nonstd_factor)
            depth=float(cat.get('std_depth_mm',0))+float(cat.get('depth_tolerance_mm',10))
            tests=[]
            if cat.get('std_depth_mm'): tests.append(f'K{dr}>{depth}')
            if cat.get('check_height') and cat.get('std_height_ft'):
                tests.append(f'ABS(M{dr}-{cat["std_height_ft"]})>{cat.get("height_tol_ft",cfg.height_tol)}')
            detail.cell(dr,18, factor if ln.nonstd_override is True else 1 if ln.nonstd_override is False else
                        '=IF(OR('+','.join(tests)+f'),{factor},1)' if tests else 1)
        if cost is not None:
            cr=4+i
            original=cost.cell(cr,8).value
            if mixed:
                prow=price_rows[(ln.cat_key,ln.material)]
                drow=price_rows.get((ln.cat_key,ln.door_key))
                cost.cell(cr,7,f'={SHEET_P}!D{prow}')
                original=f'=ROUND(D{cr}*E{cr}*G{cr}*F{cr},2)'
                if drow: original+=f'+ROUND({SHEET_D}!AC{dr}*{SHEET_D}!AD{dr}*{SHEET_P}!D{drow}*F{cr},2)'
                cost.cell(cr,8,original)
            if costrefs: cost.cell(cr,8,original+'+SUM('+','.join(costrefs)+')')
        if mixed:
            from core.kitchen import projection_rules
            rule=projection_rules(cfg.for_category(ln.cat_key))
            minimum=rule.get('min_height_ft',0) if ln.min_height_enabled else 0
            std=rule.get('std_depth_mm',600)
            for col,label in [(29,'Door area ft² / 柜门投影'),(30,'Door depth factor / 门深系数'),(31,'Door RM/ft² / 柜门单价')]:
                detail.cell(4,col,label);detail.column_dimensions[get_column_letter(col)].width=23
            detail.cell(dr,15,f'=MAX(M{dr},{minimum})')
            detail.cell(dr,16,f'=I{dr}/{cfg.mm_per_ft if ln.unit=="ft" else 1000}')
            if ln.unit=='ft':
                detail.cell(dr,19).number_format='#,##0.00######'
                wb[SHEET_P].cell(price_rows[(ln.cat_key,ln.material)],8).number_format='#,##0.00######'
                if kind=='internal':wb[SHEET_P].cell(price_rows[(ln.cat_key,ln.material)],4).number_format='#,##0.00######'
            detail.cell(dr,29,f'=L{dr}*O{dr}')
            detail.cell(dr,30,f'=MAX(1,K{dr}/{std})' if std and not rule.get('ignore_depth') else 1)
            prow=price_rows[(ln.cat_key,ln.material)]
            drow=price_rows.get((ln.cat_key,ln.door_key))
            detail.cell(dr,19,ln.unit_price if ln.price_override else f'={SHEET_P}!H{prow}')
            detail.cell(dr,31,f'={SHEET_P}!H{drow}' if drow else 0)
            detail.cell(dr,25,f'=ROUND(P{dr}*R{dr}*S{dr}*T{dr},2)')
            detail.cell(dr,26,f'=ROUND(AC{dr}*AD{dr}*AE{dr}*T{dr},2)')
            detail.cell(dr,22,f'=Y{dr}+Z{dr}+U{dr}')
    addon.freeze_panes='C2';addon.auto_filter.ref=addon.dimensions
    for c in range(1,len(heads)+1): addon.column_dimensions[get_column_letter(c)].width=18 if c!=2 else 38
    if kind=='customer': addon.sheet_state='hidden'
    else:
        if not summary['cost_complete']:
            for col in (10,11):cost.cell(4+len(lines),col,'Missing cost / 成本待补')
        cost['A2']='Material and add-on costs only; installation, transport and after-sales are in ProjectCosts. / 本表为材料及增项成本，项目总成本见项目成本表。'
        for c,text in [(13,'Actual area m² / 实际投影'),(14,'Estimated m³ / 估算体积'),(15,'Cost status / 成本状态')]:
            cost.cell(3,c,text);cost.column_dimensions[get_column_letter(c)].width=23
        pc=wb.create_sheet('项目成本_ProjectCosts',1)
        pc['A1']='Project cost & estimated profit / 项目成本与预计利润'
        pc.merge_cells('A1:D1');pc.row_dimensions[1].height=30
        p=summary['cost_policy']
        rows={4:('Installation RM/m² / 安装单价',p['installation_per_m2']),
              5:('Volume m³/m² / 投影体积系数',p['volume_per_m2']),
              6:('Transport RM/m³ / 运输单价',p['transport_per_m3']),
              7:('After-sales rate / 售后比例',p['after_sales_rate']),
              9:('Actual projection m² / 实际总投影',f'=SUM({SHEET_C}!M4:M{max(4,len(lines)+3)})'),
              10:('Estimated volume m³ / 估算运输体积','=B9*B5'),
              11:('Installation / 总安装费','=ROUND(B9*B4,2)'),
              12:('Transport / 总运输费','=ROUND(B10*B6,2)'),
              13:('Net sales before tax / 折后未税成交额',f'={SHEET_D}!V{sum_rows["net"]}'),
              14:('After-sales reserve / 售后预留成本','=ROUND(B13*B7,2)'),
              15:('Materials & add-ons / 材料及增项成本',f'=SUM({SHEET_C}!H4:H{max(4,len(lines)+3)})'),
              16:('Total project cost / 项目总成本','=SUM(B11:B12)+B14+B15'),
              17:('Estimated profit / 项目预计利润','=B13-B16' if summary['cost_complete'] else 'Missing cost / 成本待补'),
              18:('Estimated margin / 项目预计利润率','=IF(B13=0,0,B17/B13)' if summary['cost_complete'] else 'Missing cost / 成本待补')}
        for row,(label,value) in rows.items():
            pc.cell(row,1,label);pc.cell(row,2,value)
            pc.cell(row,2).number_format='0.00%' if row in (7,18) else '#,##0.0000' if row in (5,9,10) else '#,##0.00'
        pc['A20']='Area uses actual cabinet width × height × quantity; no minimum-height/depth factor. Volume is an estimate, not L×H×D.'
        pc['A21']='面积按柜体实际宽×高×数量，不重复计门、不套加价系数；体积为运输估算，不是长高深实体体积。'
        for row in (20,21):pc.merge_cells(start_row=row,start_column=1,end_row=row,end_column=4);pc.row_dimensions[row].height=30
        pc.column_dimensions['A'].width=53;pc.column_dimensions['B'].width=24;pc.column_dimensions['C'].width=15;pc.column_dimensions['D'].width=15
        pc.freeze_panes='B9';pc.sheet_view.showGridLines=False
        for row in (4,5,6,7): pc.cell(row,2).fill=PatternFill('solid',fgColor='EAF2FF')
        for row in (16,17):
            for col in (1,2):pc.cell(row,col).font=Font(bold=True,size=12)
        pc.page_setup.orientation='portrait';pc.page_setup.fitToWidth=1
    for sheet in [addon]+([pc] if kind=='internal' else []):
        for row in sheet:
            for cell in row:cell.alignment=Alignment(vertical='center',wrap_text=True)
        sheet.row_dimensions[1].height=32
