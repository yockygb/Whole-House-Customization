"""Populate cached results while preserving every editable Excel formula."""
from pathlib import Path
import zipfile,xml.etree.ElementTree as ET
from openpyxl import load_workbook
from export.formula_eval import Evaluator
NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
def recalculate(path):
    wb=load_workbook(path,rich_text=False)
    ev=Evaluator(wb);cache={}
    for sheet in wb:
        vals={}
        for row in sheet:
            for c in row:
                if c.data_type=='f':vals[c.coordinate]=ev.cell(sheet.title,c.coordinate)
        cache[sheet.title]=vals
    tmp=Path(path).with_suffix('.recalc.tmp')
    with zipfile.ZipFile(path) as src,zipfile.ZipFile(tmp,'w',zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            data=src.read(info.filename)
            if info.filename.startswith('xl/worksheets/sheet') and info.filename.endswith('.xml'):
                index=int(info.filename.split('sheet')[-1].split('.')[0])-1
                values=cache[wb.sheetnames[index]]
                tree=ET.fromstring(data)
                for c in tree.iter('{'+NS+'}c'):
                    if c.get('r') not in values:continue
                    value=values[c.get('r')]
                    old=c.find('{'+NS+'}v')
                    if old is None:old=ET.SubElement(c,'{'+NS+'}v')
                    if isinstance(value,str):c.set('t','str');old.text=value
                    elif isinstance(value,bool):c.set('t','b');old.text='1' if value else '0'
                    elif value is None:c.attrib.pop('t',None);old.text='0'
                    else:c.attrib.pop('t',None);old.text=str(value)
                data=ET.tostring(tree,encoding='utf-8',xml_declaration=True)
            dst.writestr(info,data)
    tmp.replace(path)
    return sum(len(v) for v in cache.values())
