"""Reference layout fidelity and English-first workbook typography."""
from copy import copy
from pathlib import Path
import re
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Font, Alignment
from openpyxl.worksheet.pagebreak import Break
from core.config import BASE_DIR

def finalize_workbook(wb,cfg,quote,kind,lines=None):
    ws=wb.worksheets[0]
    ref=load_workbook(Path(BASE_DIR)/'templates'/'quotation-reference.xlsx').worksheets[0]
    # Keep the old template's geometry, stationery positions and print settings.
    for key,dim in ref.column_dimensions.items():ws.column_dimensions[key].width=dim.width
    for r in range(1,9):ws.row_dimensions[r].height=ref.row_dimensions[r].height
    for r in range(1,7):
        for c in range(1,11):
            src=ref.cell(r,c);dst=ws.cell(r,c)
            if src.has_style:dst._style=copy(src._style)
    ws.page_margins=copy(ref.page_margins);ws.page_setup=copy(ref.page_setup)
    ws.print_title_rows=ref.print_title_rows
    template=cfg.raw.get('template',{})
    for col,width in template.get('column_widths',{}).items():
        if col in 'ABCDEFGHIJ' and 3<=float(width)<=70:ws.column_dimensions[col].width=float(width)
    # Keep the exact ten English headings from the old sheet; Chinese is secondary.
    from export.xlsx import _bi
    labels=[('NO.','序号'),('Product','产品名称'),('Picture','图片'),('Description','描述'),
            ('Dimension(L*H*D)mm','长*高*深 (mm)'),('Projected area','投影面积'),
            ('Unit','单位'),('QTY','数量'),('Price','单价'),('Remark','备注')]
    for c,(en,cn) in enumerate(labels,1):ws.cell(8,c).value=_bi(en,cn,en_size=9,en_bold=True)
    if quote.get('note_en') or quote.get('note') or quote.get('terms_en') or quote.get('terms'):
        r=ws.max_row+2
        ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=10)
        ws.cell(r,1).value=_bi('Project notes & terms','项目备注与条款',en_bold=True)
        for en,cn in ((quote.get('note_en',''),quote.get('note','')),(quote.get('terms_en',''),quote.get('terms',''))):
            if en or cn:
                r+=1;ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=10)
                ws.cell(r,1).value=_bi(en,cn);ws.cell(r,1).alignment=Alignment(wrap_text=True,vertical='top')
                ws.row_dimensions[r].height=max(30,14*((len(en)+len(cn)*2)//130+2))
    # Translation lookup is drawn from the user's independently editable data.
    lookup={}
    for section in ('categories','rooms','groups'):
        for entry in cfg.raw.get(section,[]):
            if entry.get('name') and entry.get('name_en'):lookup[entry['name']]=entry['name_en']
    for entries in cfg.prices.values():
        for entry in entries.values():
            if entry.get('label_en'):lookup[entry.get('label','')]=entry['label_en']
    for e in cfg.extras_items():
        if e.get('name_en'):lookup[e.get('name','')]=e['name_en']
        if e.get('note_en'):lookup[e.get('note','')]=e['note_en']
    for ln in lines or []:
        lookup[ln.material_label]=ln.material_label_en
    from export.descriptions import remark
    from export.xlsx import _rt_lines
    by_index={ln.idx:ln for ln in lines or []}
    for row in ws:
        index=row[0].value
        if isinstance(index,int) and index in by_index and row[1].value is not None:
            en,cn=remark(by_index[index],cfg)
            ws.cell(row[0].row,10).value=_rt_lines([(en,False),(cn,True)],en_size=8)
    # Render company-supplied bilingual captions without dropping either language.
    co=cfg.company()
    for row,key,cn_key in ((1,'name','name_cn'),(2,'address','address_cn')):
        if co.get(cn_key):
            ws.cell(row,3).value=_bi(co.get(key,''),co[cn_key],en_size=9,en_bold=(row==1))
            ws.row_dimensions[row].height=max(24,ws.row_dimensions[row].height or 12)
    if co.get('greeting_cn'):
        ws['A6']=_bi(co.get('greeting',''),co['greeting_cn'],en_size=9)
        ws.row_dimensions[6].height=30
    for sheet in wb:
        for row in sheet:
            for cell in row:
                v=cell.value
                if isinstance(v,str) and not v.startswith('=') and v in lookup:
                    cell.value=lookup[v]+'\n'+v
                elif isinstance(v,str) and not v.startswith('=') and '\n' in v:
                    chunks=v.split('\n')
                    # Internal headers and source labels inherit the same English-first order.
                    if re.search('[\u4e00-\u9fff]',chunks[0]) and any(re.search('[A-Za-z]',x) and not re.search('[\u4e00-\u9fff]',x) for x in chunks[1:]):
                        cell.value='\n'.join([x for x in chunks if not re.search('[\u4e00-\u9fff]',x)]+[x for x in chunks if re.search('[\u4e00-\u9fff]',x)])
        if kind=='customer' and 'PriceList' in sheet.title:
            # Hidden is not removed: erase supply and market values as well as hiding those columns.
            for row in sheet.iter_rows(min_col=4,max_col=7):
                for cell in row:
                    if not isinstance(cell,MergedCell):cell.value=None
    detail=next((s for s in wb if 'Calculation' in s.title),None)
    if detail:
        if kind=='customer':detail['A1']='Calculation detail / 计算明细'
        for n,ln in enumerate(lines or [],5):
            if ln.room_note or ln.room_note_en:detail.cell(n,4).value=(ln.room_note_en or 'English translation needed')+'\n'+ln.room_note
            if ln.extras_list or getattr(ln, 'led_type', 'none') != 'none':
                en,cn=remark(ln,cfg)
                detail.cell(n,23).value=en+'\n'+cn
    wb.calculation.fullCalcOnLoad=True
    wb.calculation.forceFullCalc=True

