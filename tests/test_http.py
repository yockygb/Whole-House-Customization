"""Start the real app with disposable data and exercise its public workflow."""
from pathlib import Path
import io
import json
import os
import socket
import subprocess
import sys
import tempfile
import time
import unittest
import urllib.request

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]


class HttpTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory(prefix='whole-house-customization-http-')
        with socket.socket() as sock:
            sock.bind(('127.0.0.1', 0))
            port = sock.getsockname()[1]
        cls.base = 'http://127.0.0.1:%d' % port
        env = dict(os.environ,
                   WHOLE_HOUSE_CUSTOMIZATION_DATA_DIR=str(Path(cls.tmp.name) / 'data'),
                   WHOLE_HOUSE_CUSTOMIZATION_EXPORT_DIR=str(Path(cls.tmp.name) / 'exports'),
                   WHOLE_HOUSE_CUSTOMIZATION_NO_BROWSER='1', PYTHONIOENCODING='utf-8')
        cls.log = open(Path(cls.tmp.name) / 'server.log', 'w+', encoding='utf-8')
        cls.proc = subprocess.Popen(
            [sys.executable, str(ROOT / 'app.py'), '--no-browser', '--port', str(port)],
            cwd=ROOT, env=env, stdout=cls.log, stderr=cls.log,
            creationflags=0x08000000 if os.name == 'nt' else 0)
        try:
            for _ in range(100):
                if cls.proc.poll() is not None:
                    cls.log.seek(0)
                    raise RuntimeError(cls.log.read())
                try:
                    with urllib.request.urlopen(cls.base + '/api/identity', timeout=0.5) as response:
                        if response.status == 200:
                            return
                except OSError:
                    time.sleep(0.1)
            raise RuntimeError('Application startup timed out')
        except BaseException:
            cls.tearDownClass()
            raise

    @classmethod
    def tearDownClass(cls):
        cls.proc.terminate()
        cls.proc.wait(timeout=10)
        cls.log.close()
        cls.tmp.cleanup()

    def request_json(self, path, data=None, method=None):
        request = urllib.request.Request(
            self.base + path,
            data=json.dumps(data).encode('utf-8') if data is not None else None,
            headers={'Content-Type': 'application/json'}, method=method)
        with urllib.request.urlopen(request, timeout=15) as response:
            return json.load(response)

    def test_project_name_in_identity_and_page(self):
        identity = self.request_json('/api/identity')
        self.assertEqual(identity['application'], 'whole-house-customization-independent-v3')
        self.assertEqual(identity['version'], '3.0.2')
        with urllib.request.urlopen(self.base + '/', timeout=5) as response:
            html = response.read().decode('utf-8')
        self.assertIn('<title>Whole-House Customization 报价系统</title>', html)
        self.assertIn('<strong>Whole-House Customization</strong>', html)

    def test_bootstrap_and_empty_database(self):
        self.assertTrue(self.request_json('/api/health')['ok'])
        self.assertEqual(self.request_json('/api/quotes')['quotes'], [])
        self.assertTrue(self.request_json('/api/bootstrap'))
        with urllib.request.urlopen(self.base + '/', timeout=5) as response:
            self.assertIn(b'id="depth"', response.read())

    def test_custom_category_material_can_be_created_edited_and_deleted(self):
        config = self.request_json('/api/config')['config']
        key = 'test_category'
        config['categories'].append({
            'key': key, 'name': '测试品类', 'name_en': 'Test Category',
            'group': 'cabinet', 'measure': 'area', 'std_depth_mm': 600,
            'min_height_ft': 0, 'default_height_mm': 2400,
        })
        config['prices'][key] = {
            'standard': {'label': '标准材料', 'label_en': 'Standard',
                         'part': 'body', 'unit': 'ft2', 'supply': 0,
                         'market': [0, 0]}
        }
        self.request_json('/api/config', {'config': config}, method='PUT')
        created = self.request_json('/api/materials', {
            'cat_key': key, 'mat_key': 'test_board', 'part': 'body',
            'name': '测试板材', 'name_en': 'Test Board', 'unit': 'ft2',
            'quote_price': 0, 'cost_price': 0,
        })['material']
        self.assertEqual(created['cat_key'], key)
        self.assertIn(key, [c['key'] for c in self.request_json('/api/bootstrap')['categories']])
        self.request_json('/api/materials/%d' % created['id'], {
            'name': '已修改材料', 'name_en': 'Edited Board',
        }, method='PUT')
        quote = self.request_json('/api/quotes', {
            'customer': 'Test Customer', 'project': 'Test Project',
            'items': [{'cat': key, 'material': 'test_board',
                       'length_mm': 1200, 'height_mm': 2400,
                       'depth_mm': 600, 'qty': 1, 'unit_price': 0}],
        })['quote']
        config = self.request_json('/api/config')['config']
        config['categories'] = [c for c in config['categories'] if c['key'] != key]
        config['prices'].pop(key, None)
        self.request_json('/api/config', {'config': config}, method='PUT')
        self.assertNotIn(key, [c['key'] for c in self.request_json('/api/bootstrap')['categories']])
        archived = self.request_json('/api/config')['config']
        archived_category = next(c for c in archived['categories'] if c['key'] == key)
        self.assertFalse(archived_category['enabled'])
        self.assertIn('test_board', archived['prices'][key])
        saved_quote = self.request_json('/api/quotes/%d' % quote['id'])['quote']
        self.assertEqual(saved_quote['items'][0]['cat'], key)

    def test_option_order_is_persisted_and_returned_to_quote_form(self):
        before = self.request_json('/api/bootstrap')['categories']
        original = [c['key'] for c in before]
        changed = list(reversed(original))
        result = self.request_json('/api/catalogue/order', {
            'collection': 'categories', 'order': changed,
        })
        self.assertEqual([c['key'] for c in result['bootstrap']['categories']], changed)
        self.request_json('/api/catalogue/order', {
            'collection': 'categories', 'order': original,
        })

    def test_floor_order_is_returned_to_quote_form(self):
        original = [f['code'] for f in self.request_json('/api/bootstrap')['floors']]
        changed = list(reversed(original))
        result = self.request_json('/api/catalogue/order', {
            'collection': 'floors', 'order': changed,
        })
        self.assertEqual([f['code'] for f in result['bootstrap']['floors']], changed)
        self.request_json('/api/catalogue/order', {
            'collection': 'floors', 'order': original,
        })

    def test_create_quote_and_download_customer_workbook(self):
        result = self.request_json('/api/quotes', {
            'customer': 'Demo Customer', 'project': 'Open Source Demo',
            'items': [{'cat': 'custom_cabinet', 'material': 'body_particle',
                       'length_mm': 1219.2, 'height_mm': 1219.2,
                       'depth_mm': 600, 'qty': 1, 'unit_price': 1}]})
        self.assertEqual(result['computed']['summary']['total'], 16)
        quote_id = result['quote']['id']
        exported = self.request_json('/api/quotes/%d/export?kind=customer&fmt=xlsx' % quote_id)
        with urllib.request.urlopen(self.base + exported['url'], timeout=15) as response:
            workbook = load_workbook(io.BytesIO(response.read()))
        self.assertTrue(workbook.worksheets[0]._images)
        self.assertFalse(any('CostProfit' in name for name in workbook.sheetnames))
        workbook.close()


if __name__ == '__main__':
    unittest.main(verbosity=2)
