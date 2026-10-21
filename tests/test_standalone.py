"""Independent tests; all business data stays in a temporary directory."""
import os,sys,tempfile,json,copy,unittest
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT));sys.path.insert(0,str(ROOT/'tests'))
TMP=tempfile.TemporaryDirectory(prefix='whole-house-customization-selftest-')
os.environ['WHOLE_HOUSE_CUSTOMIZATION_DATA_DIR']=str(Path(TMP.name)/'data')
os.environ['WHOLE_HOUSE_CUSTOMIZATION_EXPORT_DIR']=str(Path(TMP.name)/'exports')
from server import store
from server.independent import save_config,backup_payload,export_config,validate_config
from core import Item,compute_quote
from fixtures import DEMO
from export.xlsx import export_workbook
from openpyxl import load_workbook

class StandaloneTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):store.init_db()
    def setUp(self):
        self.c=store.connect();self.raw=copy.deepcopy(store.pricing_raw())
    def tearDown(self):
        save_config(self.c,self.raw);self.c.close()
    def test_material_name_persists_after_restart(self):
        m=store.list_materials(self.c)[0]
        store.update_material(self.c,m['id'],{'name':'独立材料','name_en':'Independent material','quote_price':1})
        store.init_db()
        found=next(x for x in store.list_materials(self.c) if x['id']==m['id'])
        self.assertEqual(found['name_en'],'Independent material');self.assertEqual(found['quote_price'],1)
        store.update_material(self.c,m['id'],{'name':m['name'],'name_en':m['name_en'],'quote_price':None})
    def test_rules_persist_after_restart(self):
        raw=copy.deepcopy(self.raw);raw['categories'][0]['std_depth_mm']=555
        save_config(self.c,raw);store.init_db();self.assertEqual(store.build_config(self.c).category('custom_cabinet')['std_depth_mm'],555)
    def test_depth_and_bilingual_notes(self):
        it=copy.deepcopy(DEMO[0]);it.update(depth_mm=720,note_en='English note',room_note_en='English location')
        q=store.create_quote(self.c,{'items':[it]});got=store.get_quote(self.c,q['id'])['items'][0]
        self.assertEqual(got['depth_mm'],720);self.assertEqual(got['note_en'],'English note');self.assertEqual(got['room_note_en'],'English location')
    def test_locked_rule_snapshot(self):
        q=store.create_quote(self.c,{'items':DEMO});locked=store.lock_quote(self.c,q['id']);before=store.compute_for(self.c,locked)['summary']['total']
        raw=copy.deepcopy(self.raw);raw['categories'][0]['std_depth_mm']=300;save_config(self.c,raw)
        self.assertEqual(store.compute_for(self.c,store.get_quote(self.c,q['id']))['summary']['total'],before)
        self.assertEqual(export_config(self.c,locked).category('custom_cabinet')['std_depth_mm'],600)
    def test_zero_material_and_line_prices(self):
        cfg=store.build_config(self.c);it=Item.from_dict({**DEMO[0],'unit_price':0,'extras':{}})
        lines,_=compute_quote(cfg,[it]);self.assertEqual(lines[0].amount,0)
    def test_customer_workbook_has_pictures_and_cached_formulas(self):
        cfg=store.build_config(self.c);lines,summary=compute_quote(cfg,[Item.from_dict(i) for i in DEMO])
        path=Path(TMP.name)/'sample.xlsx';export_workbook(str(path),cfg,{'quote_no':'TEST','items':DEMO},lines,summary,'customer')
        w=load_workbook(path);self.assertEqual(len(w.worksheets[0]._images),len(DEMO))
        self.assertFalse(any('CostProfit' in s.title for s in w))
        prices=next(s for s in w if 'PriceList' in s.title)
        self.assertTrue(all(c.value is None for row in prices.iter_rows(min_col=4,max_col=7) for c in row))
        cached=load_workbook(path,data_only=True)
        for s in w:
            for row in s:
                for c in row:
                    if c.data_type=='f':self.assertIsNotNone(cached[s.title][c.coordinate].value)
    def test_backup_contains_independent_data(self):
        b=backup_payload(self.c);self.assertEqual(b['format'],'Whole-House-Customization-Backup');self.assertIn('quote_items',b['tables']);self.assertIn('prices',b['pricing'])
    def test_invalid_configuration_rejected(self):
        raw=copy.deepcopy(self.raw);raw['meta']['mm_per_ft']=0
        with self.assertRaises(ValueError):validate_config(raw)
    def test_dimension_candidates_do_not_invent_height(self):
        from server.drawing import candidates
        d=candidates('2400 x 600 mm')[0]
        self.assertEqual(d['height_mm'],0);self.assertEqual(d['depth_mm'],600)
    def test_original_template_widths_preserved(self):
        reference=load_workbook(ROOT/'templates'/'quotation-reference.xlsx').worksheets[0]
        self.assertEqual(reference.column_dimensions['C'].width,26)
        self.assertEqual(reference.page_setup.fitToWidth,1)
    def test_excel_rules_recalculate_after_dimension_changes(self):
        from export.formula_eval import Evaluator
        cfg=store.build_config(self.c)
        for table in cfg.prices.values():
            for entry in table.values():entry['market']=[1,1]
        it={**DEMO[0],'extras':{},'height_mm':2400}
        lines,summary=compute_quote(cfg,[Item.from_dict(it)])
        path=Path(TMP.name)/'mutable.xlsx';export_workbook(str(path),cfg,{'quote_no':'MUTABLE'},lines,summary,'internal')
        w=load_workbook(path);detail=w['明细_Calculation'];detail['J5']=500
        ev=Evaluator(w);self.assertAlmostEqual(ev.cell(detail.title,'O5'),3.3)
        expected=compute_quote(cfg,[Item.from_dict({**it,'height_mm':500})])[0][0].amount
        self.assertAlmostEqual(ev.cell(detail.title,'V5'),expected,places=2)
        it={**DEMO[4],'extras':{},'height_mm':1000,'nonstd':None}
        lines,summary=compute_quote(cfg,[Item.from_dict(it)])
        export_workbook(str(path),cfg,{'quote_no':'MUTABLE'},lines,summary,'internal')
        w=load_workbook(path);w['明细_Calculation']['J5']=914.4
        self.assertEqual(Evaluator(w).cell('明细_Calculation','R5'),1)

if __name__=='__main__':
    try:unittest.main(verbosity=2)
    finally:
        import gc
        gc.collect()
        TMP.cleanup()
